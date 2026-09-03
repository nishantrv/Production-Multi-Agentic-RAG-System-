import logfire
from docx import Document
from pptx import Presentation

def parse_office(file_path: str):
    """
    Parses Office documents (.docx, .pptx) locally.
    """
    with logfire.span("📄 Office Document Parsing", filename=file_path):
        try:
            if file_path.lower().endswith(".docx"):
                document = Document(file_path)
                text_parts = [paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    text_parts.extend(
                        "\t".join(cell.text for cell in row.cells)
                        for row in table.rows
                    )
            elif file_path.lower().endswith(".pptx"):
                presentation = Presentation(file_path)
                text_parts = [
                    shape.text
                    for slide in presentation.slides
                    for shape in slide.shapes
                    if hasattr(shape, "text")
                ]
            else:
                raise ValueError(f"Unsupported Office file type: {file_path}")

            full_text = "\n".join(text_parts)
            
            if not full_text.strip():
                logfire.warning(f"⚠️ Office parser returned empty text for {file_path}")
            else:
                logfire.info(f"✅ Successfully parsed {len(full_text)} characters")

            return full_text
        except Exception as e:
            logfire.error(f"❌ Office Parse Failed: {e}")
            raise e